import sys
import subprocess
import os

try:
    from docx import Document
except ImportError:
    print("Installing python-docx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

# Create document
doc = Document()

# Title
doc.add_heading('Graphics API Practical - Task Implementation', 0)

# Section 1
doc.add_heading('Unique Implementation Details', level=1)
doc.add_paragraph('For this practical, I have implemented a "Catch the Stars" 2D arcade game. The game demonstrates the full graphics pipeline by rendering a player-controlled basket catching falling objects (stars) while avoiding hazards (asteroids). The unique implementations chosen include:')

doc.add_paragraph('1. Dynamic Speed Scaling (Difficulty): The game becomes progressively harder. As the player catches more stars and their score increases, the level goes up, and the falling speed of all objects increases on the screen.')
doc.add_paragraph('2. Synthesized Audio: Real-time audio generation using the Web Audio API (Oscillators). A high-pitched sine wave plays when a star is caught, and a low-pitched sawtooth wave plays when an asteroid is hit.')
doc.add_paragraph('3. Procedural Geometry: The asteroids are not perfect circles; they are dynamically generated polygons with slightly randomized radii at each vertex, giving them a jagged, rock-like appearance every time they spawn.')
doc.add_paragraph('4. Layered Rendering Effects: Uses HTML5 Canvas shadow properties for glowing effects on the player and semi-transparent overlay circles inside the stars for visual depth.')

# Section 2
doc.add_heading('Graphics Pipeline Implementation Steps', level=1)
doc.add_paragraph('The game strictly follows and demonstrates the three main stages of the graphics pipeline, which are heavily commented in the script.js file.')

doc.add_heading('1. Application Stage', level=2)
doc.add_paragraph('In this stage, the logic and rules are defined. Variables for score, level, and player coordinates are initialized. The types of objects (Stars and Asteroids) are defined with properties like coordinates (x, y), speed, rotation, color, and size. User input (keyboard events) is also captured here to update the player\'s velocity vector (dx). The rules dictate that the game ends or levels up based on the score.')

doc.add_heading('2. Geometry Stage', level=2)
doc.add_paragraph('During the game loop, the coordinates of all objects are transformed. The player object is translated horizontally based on user input. The falling objects are translated vertically based on their calculated speed and have their rotation angles updated. Finally, bounding-box (AABB) and circular collision checks are calculated geometrically to determine if objects intersect, simulating a physical interaction in 2D space.')

doc.add_heading('3. Rasterization Stage', level=2)
doc.add_paragraph('In the final stage, mathematical geometry is converted into pixels on the screen. The screen grid is cleared first. The player is scan-converted into a solid rectangle with a shadow blur effect. The Canvas API translates and rotates the rendering context to match the geometry stage coordinates. Primitive assembly connects the vertices of stars and asteroids into paths, which are then filled with colors and stroked with outlines to map them onto the final pixel grid. Overlapping transparency (RGBA) is used for paused states and visual details.')

# Save docx
doc.save('task.docx')
print("Successfully generated task.docx")
